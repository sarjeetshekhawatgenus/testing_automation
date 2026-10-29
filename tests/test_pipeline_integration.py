"""Integration tests for the transformation pipeline across all nine operations."""

from __future__ import annotations

import os
import tempfile

import pandas as pd
import pytest
from openpyxl import Workbook

from master_update_automation.config import (
    AuditStatus,
    Category,
    OPERATIONS,
    Operation,
)
from master_update_automation.excel_writer import write_output_workbook
from master_update_automation.pipeline import run_pipeline
from tests.conftest import (
    AMBIG_DT_A,
    AMBIG_DT_B,
    CIRCLE_ID,
    DISCOM_ID,
    DIV_ID,
    DT_ID,
    ESC_ID,
    FEEDER_ID,
    GENUS_HTCT,
    GENUS_LTCT,
    GENUS_WC,
    LEC_ID,
    MAIN_METER_ID,
    NSC_ID,
    SECURE_LTCT,
    SUBDIV_ID,
    TOD_ID,
    ZONE_ID,
    make_template_workbook,
    operation,
    reference_frame,
)

# ── helpers ──────────────────────────────────────────────────────────────────


def _tmpdir():
    return tempfile.mkdtemp()


def _template_path(tmpdir: str) -> str:
    path = os.path.join(tmpdir, "template.xlsx")
    make_template_workbook(path)
    return path


def _hierarchy_csv(tmpdir: str) -> str:
    path = os.path.join(tmpdir, "hierarchy.csv")
    df = pd.DataFrame([
        {
            "Discom ID": DISCOM_ID, "Discom Name": "JVVNL", "Discom Code": "GJV96",
            "Zone ID": ZONE_ID, "Zone Name": "Jaipur", "Zone Code": "JPR",
            "Circle ID": CIRCLE_ID, "Circle Name": "KOTPUTLI", "Circle Code": "2117000",
            "Division ID": DIV_ID, "Division Name": "XEN_OM_KOTPUTALI", "Division Code": "2117200",
            "Subdivision ID": SUBDIV_ID, "Subdivision Name": "AEN_OM_VIRAT NAGAR",
            "Subdivision Code": "2117240",
            "Substation ID": "88888888-8888-8888-8888-888888888888",
            "Substation Name": "AANTELA", "Substation Code": "1099",
            "Feeder ID": FEEDER_ID, "Feeder Name": "ODA CRUSHER", "Feeder Code": "1085140",
            "DT ID": DT_ID, "DT Name": "GOVINDAM MINERALS", "DT Code": "1085140006",
        },
    ])
    df.to_csv(path, index=False, encoding="utf-8-sig")
    return path


def _source_excel(tmpdir: str, rows: list[dict], sheet: str = "Sheet1") -> str:
    path = os.path.join(tmpdir, "source.xlsx")
    df = pd.DataFrame(rows)
    df.to_excel(path, index=False, sheet_name=sheet, engine="openpyxl")
    return path


def _base_source_row(**overrides) -> dict:
    defaults = {
        "Consumer Number": "C001",
        "Account Number": "A001",
        "Consumer UAN Suffix": "UAN001",
        "Consumer Name": "Test Consumer",
        "Father Name": "Father",
        "Consumer Address": "Address",
        "Mobile Number": "9876543210",
        "DT Code": "1085140006",
        "Meter Category": "Main Meter",
        "Consumer Status": "Regular",
        "Meter Phase": "3-Phase",
        "Meter Serial No": "MSN001",
        "Meter Make": "Genus",
        "Meter Status": "OK",
        "Geographical Area": "Rural",
        "Is Consumer Available In Master?": "Yes",
        "Consumer Connection Type": "ESC",
        "Consumer Identification": "ESC",
        "Is Old Meter Available?": "Yes",
        "Sanctioned Load Unit": "HP",
        "Connected Load": "10",
        "Sanctioned Load": "10",
    }
    defaults.update(overrides)
    return defaults


def _run(op_src: str, op_dst: str, source_rows: list[dict], tmpdir: str | None = None):
    d = tmpdir or _tmpdir()
    tpl = _template_path(d)
    hier = _hierarchy_csv(d)
    src = _source_excel(d, source_rows)
    op = operation(op_src, op_dst)
    result = run_pipeline(
        template_path=tpl,
        target_sheet="HTCT Consumer Master",
        operation=op,
        hierarchy_path=hier,
        source_path=src,
        source_sheet="Sheet1",
    )
    return result


def _run_with_hierarchy(op_src: str, op_dst: str, source_rows: list[dict], hierarchy_df: pd.DataFrame):
    d = _tmpdir()
    tpl = _template_path(d)
    hier = os.path.join(d, "hierarchy.csv")
    hierarchy_df.to_csv(hier, index=False, encoding="utf-8-sig")
    src = _source_excel(d, source_rows)
    return run_pipeline(
        template_path=tpl,
        target_sheet="HTCT Consumer Master",
        operation=operation(op_src, op_dst),
        hierarchy_path=hier,
        source_path=src,
        source_sheet="Sheet1",
    )


# ── all 9 operations smoke tests ─────────────────────────────────────────────


@pytest.mark.parametrize("op_id", list(range(1, 10)))
def test_pipeline_all_nine_ops_produce_records(op_id):
    """Every operation must produce output records from a valid source row."""
    op = OPERATIONS[op_id]
    ident = "ESC" if op.source == Category.WC else "TOD"
    make = "Genus"
    result = _run(
        op.source.value,
        op.destination.value,
        [_base_source_row(**{"Consumer Identification": ident, "Meter Make": make})],
    )
    assert len(result.records) == 1
    assert result.audit.total_records == 1


# ── UUID formatting ──────────────────────────────────────────────────────────


def test_all_uuids_in_output_are_braced():
    """Every non-empty UUID field in the output must be {uuid} format."""
    result = _run("WC", "HTCT", [_base_source_row()])
    import re
    uuid_pat = re.compile(r"^\{[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\}$")
    from master_update_automation.config import UUID_FIELDS
    for rec in result.records:
        for key, val in rec.items():
            if key.lower().strip() in UUID_FIELDS and val and val.strip():
                assert uuid_pat.match(val.strip()), (
                    f"Field '{key}' has non-canonical UUID: {val}"
                )


def test_no_double_braces_in_output():
    """Output must never contain {{uuid}}."""
    result = _run("WC", "HTCT", [_base_source_row()])
    for rec in result.records:
        for val in rec.values():
            if val and isinstance(val, str):
                assert "{{" not in val, f"Double braces found: {val}"
                assert "}}" not in val, f"Double braces found: {val}"


# ── ESC → TOD override (Rule A) ─────────────────────────────────────────────


class TestRuleA_ESC_TO_TOD:
    """WC → HTCT or WC → LTCT with ESC must resolve to TOD UUID."""

    def test_wc_to_htct_esc_becomes_tod(self):
        result = _run("WC", "HTCT", [_base_source_row(**{"Consumer Identification": "ESC"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{TOD_ID}}}"
        assert ESC_ID not in rec.get("consumer_identy", "")

    def test_wc_to_ltct_esc_becomes_tod(self):
        result = _run("WC", "LTCT", [_base_source_row(**{"Consumer Identification": "ESC"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{TOD_ID}}}"

    def test_wc_to_htct_non_esc_does_not_become_tod(self):
        result = _run("WC", "HTCT", [_base_source_row(**{"Consumer Identification": "NSC"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{NSC_ID}}}"
        assert TOD_ID not in rec.get("consumer_identy", "")

    def test_wc_to_ltct_non_esc_does_not_become_tod(self):
        result = _run("WC", "LTCT", [_base_source_row(**{"Consumer Identification": "LEC"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{LEC_ID}}}"


# ── Rule B: HTCT/LTCT → WC ──────────────────────────────────────────────────


class TestRuleB_ToWC:
    """HTCT/LTCT → WC uses the source identification UUID, not TOD."""

    def test_htct_to_wc_preserves_source_ident(self):
        result = _run("HTCT", "WC", [_base_source_row(**{"Consumer Identification": "LEC"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{LEC_ID}}}"

    def test_ltct_to_wc_preserves_source_ident(self):
        result = _run("LTCT", "WC", [_base_source_row(**{"Consumer Identification": "NSC"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{NSC_ID}}}"

    def test_htct_to_wc_esc_stays_esc(self):
        """Even with ESC, HTCT→WC should NOT apply the ESC→TOD override."""
        result = _run("HTCT", "WC", [_base_source_row(**{"Consumer Identification": "ESC"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{ESC_ID}}}"
        assert TOD_ID not in rec.get("consumer_identy", "")

    def test_ltct_to_wc_with_different_idents(self):
        rows = [
            _base_source_row(**{"Consumer Number": "C001", "Account Number": "A001",
                                "Consumer Identification": "ESC"}),
            _base_source_row(**{"Consumer Number": "C002", "Account Number": "A002",
                                "Consumer Identification": "TOD"}),
            _base_source_row(**{"Consumer Number": "C003", "Account Number": "A003",
                                "Consumer Identification": "NSC"}),
        ]
        result = _run("LTCT", "WC", rows)
        assert len(result.records) == 3
        assert result.records[0]["consumer_identy"].strip() == f"{{{ESC_ID}}}"
        assert result.records[1]["consumer_identy"].strip() == f"{{{TOD_ID}}}"
        assert result.records[2]["consumer_identy"].strip() == f"{{{NSC_ID}}}"


# ── Rule C: Same-category ────────────────────────────────────────────────────


class TestRuleC_SameCategory:
    """Same-category updates must NOT apply the ESC→TOD override."""

    def test_wc_to_wc_esc_stays_esc(self):
        result = _run("WC", "WC", [_base_source_row(**{"Consumer Identification": "ESC"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{ESC_ID}}}"

    def test_ltct_to_ltct_standard_lookup(self):
        result = _run("LTCT", "LTCT", [_base_source_row(**{"Consumer Identification": "TOD"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{TOD_ID}}}"

    def test_htct_to_htct_standard_lookup(self):
        result = _run("HTCT", "HTCT", [_base_source_row(**{"Consumer Identification": "NSC"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{NSC_ID}}}"

    def test_same_category_still_transforms_lookups(self):
        """Even WC→WC must transform reference fields, not just copy."""
        result = _run("WC", "WC", [_base_source_row()])
        rec = result.records[0]
        # meter_category should be the Main Meter UUID, not plain text
        mc = rec.get("meter_category", "").strip()
        assert mc == f"{{{MAIN_METER_ID}}}", f"meter_category should be Main Meter UUID, got: {mc}"


# ── Rule D: LTCT ↔ HTCT ─────────────────────────────────────────────────────


class TestRuleD_CrossCategory:
    """LTCT↔HTCT uses destination meter make and standard identification lookup."""

    def test_ltct_to_htct_meter_make(self):
        result = _run("LTCT", "HTCT", [_base_source_row(**{"Meter Make": "Genus"})])
        rec = result.records[0]
        mm = rec.get("meter_make", "").strip()
        assert mm == f"{{{GENUS_HTCT}}}", f"Expected HTCT Genus UUID, got: {mm}"

    def test_htct_to_ltct_meter_make(self):
        result = _run("HTCT", "LTCT", [_base_source_row(**{"Meter Make": "Genus"})])
        rec = result.records[0]
        mm = rec.get("meter_make", "").strip()
        assert mm == f"{{{GENUS_LTCT}}}", f"Expected LTCT Genus UUID, got: {mm}"

    def test_ltct_to_htct_esc_not_overridden(self):
        """LTCT→HTCT must NOT apply Rule A override."""
        result = _run("LTCT", "HTCT", [_base_source_row(**{"Consumer Identification": "ESC"})])
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{ESC_ID}}}"


# ── Meter make ───────────────────────────────────────────────────────────────


class TestMeterMakeCategories:
    """Meter make UUID must correspond to the destination category."""

    def test_wc_to_htct_uses_htct_make(self):
        result = _run("WC", "HTCT", [_base_source_row(**{"Meter Make": "Genus"})])
        assert result.records[0].get("meter_make", "").strip() == f"{{{GENUS_HTCT}}}"

    def test_wc_to_ltct_uses_ltct_make(self):
        result = _run("WC", "LTCT", [_base_source_row(**{"Meter Make": "Genus"})])
        assert result.records[0].get("meter_make", "").strip() == f"{{{GENUS_LTCT}}}"

    def test_htct_to_wc_uses_wc_make(self):
        result = _run("HTCT", "WC", [_base_source_row(**{"Meter Make": "Genus"})])
        assert result.records[0].get("meter_make", "").strip() == f"{{{GENUS_WC}}}"

    def test_missing_make_flagged(self):
        result = _run("WC", "HTCT", [_base_source_row(**{"Meter Make": "NONEXISTENT"})])
        entry = result.audit.entries[0]
        assert any("meter_make" in fl for fl in entry.failed_lookups)


# ── Hierarchy resolution ─────────────────────────────────────────────────────


class TestHierarchyResolution:
    def test_hierarchy_uuids_populated(self):
        result = _run("WC", "HTCT", [_base_source_row()])
        rec = result.records[0]
        assert rec.get("dt_name", "").strip() == f"{{{DT_ID}}}"
        assert rec.get("feeder_name", "").strip() == f"{{{FEEDER_ID}}}"
        assert rec.get("subdivision_name", "").strip() == f"{{{SUBDIV_ID}}}"
        assert rec.get("division_name", "").strip() == f"{{{DIV_ID}}}"
        assert rec.get("circle_name", "").strip() == f"{{{CIRCLE_ID}}}"
        assert rec.get("zone_name", "").strip() == f"{{{ZONE_ID}}}"
        assert rec.get("discom_name", "").strip() == f"{{{DISCOM_ID}}}"

    def test_missing_hierarchy_flagged_not_invented(self):
        result = _run("WC", "HTCT", [_base_source_row(**{"DT Code": "DOESNOTEXIST"})])
        entry = result.audit.entries[0]
        assert entry.errors  # hierarchy errors expected
        rec = result.records[0]
        # DT UUID should be empty, not invented
        assert rec.get("dt_name", "").strip() == "" or "DOESNOTEXIST" not in rec.get("dt_name", "")

    def test_feeder_only_source_does_not_inherit_arbitrary_dt(self, hierarchy_df):
        row = _base_source_row()
        row.pop("DT Code")
        row["Feeder Code"] = "1085140"
        result = _run_with_hierarchy("WC", "HTCT", [row], hierarchy_df)
        rec = result.records[0]
        assert rec["feeder_name"] == f"{{{FEEDER_ID}}}"
        assert rec["dt_name"] == ""
        assert rec["dt_code"] == ""
        dt_field = next(
            field for field in result.audit.entries[0].field_resolutions
            if field.target_field == "dt_name"
        )
        assert dt_field.status == "UNRESOLVED"
        assert "Feeder Code" in dt_field.source_field


def test_mandatory_blank_source_is_reported_and_not_successful():
    result = _run("WC", "HTCT", [_base_source_row(**{"Consumer Name": ""})])
    assert result.records[0]["consumer_name"] == ""
    assert not result.ready_for_output
    field = next(
        item for item in result.audit.entries[0].field_resolutions
        if item.target_field == "consumer_name"
    )
    assert field.status == "UNRESOLVED"
    assert field.mandatory
    assert field.source_field == "Consumer Name"
    report = result.audit.get_field_error_report()
    assert "Consumer Name" in set(report["Source Field"])


def test_absent_mandatory_source_column_is_not_substituted():
    row = _base_source_row()
    row.pop("Consumer Name")
    result = _run("WC", "HTCT", [row])
    assert result.records[0]["consumer_name"] == ""
    assert not result.ready_for_output
    field = next(
        item for item in result.audit.entries[0].field_resolutions
        if item.target_field == "consumer_name"
    )
    assert field.status == "UNRESOLVED"
    assert field.source_field == ""
    assert field.mandatory


def test_invalid_direct_uuid_is_left_blank_and_marked_invalid():
    row = _base_source_row(**{"Response ID": "NOT-A-UUID"})
    result = _run("WC", "HTCT", [row])
    assert result.records[0]["id"] == ""
    field = next(
        item for item in result.audit.entries[0].field_resolutions
        if item.target_field == "id"
    )
    assert field.status == "INVALID"
    assert field.source_field == "Response ID"


def test_conflicting_source_aliases_do_not_choose_first_column():
    row = _base_source_row(**{"Consumer UAN": "FIRST", "Consumer UAN Suffix": "SECOND"})
    result = _run("WC", "HTCT", [row])
    assert result.records[0]["consumer_uan_suffix"] == ""
    assert not result.ready_for_output
    field = next(
        item for item in result.audit.entries[0].field_resolutions
        if item.target_field == "consumer_uan_suffix"
    )
    assert field.status == "AMBIGUOUS"
    assert set(field.source_field.split("; ")) == {"Consumer UAN", "Consumer UAN Suffix"}


def test_optional_reference_without_match_is_blank_and_audited():
    result = _run(
        "WC",
        "HTCT",
        [_base_source_row(**{"Geographical Area": "UNLISTED AREA"})],
    )
    assert result.records[0]["geographical_are"] == ""
    field = next(
        item for item in result.audit.entries[0].field_resolutions
        if item.target_field == "geographical_are"
    )
    assert field.status == "UNRESOLVED"
    assert field.source_field == "Geographical Area"
    assert "not found" in field.reason.lower()


def test_several_missing_values_stay_blank_without_fallbacks():
    row = _base_source_row(**{
        "Meter Serial No": "",
        "Meter Make": "",
        "Geographical Area": "",
        "DT Code": "NO-SUCH-DT",
    })
    result = _run("WC", "HTCT", [row])
    record = result.records[0]
    assert record["meter_serial_no"] == ""
    assert record["meter_make"] == ""
    assert record["geographical_are"] == ""
    assert record["dt_name"] == ""
    field_status = {
        item.target_field: item.status
        for item in result.audit.entries[0].field_resolutions
    }
    assert field_status["meter_serial_no"] == "UNRESOLVED"
    assert field_status["meter_make"] == "UNRESOLVED"
    assert field_status["geographical_are"] == "UNRESOLVED"


def test_valid_source_values_are_preserved():
    result = _run(
        "WC",
        "HTCT",
        [_base_source_row(**{"Consumer Number": "C-TRACE-01", "Consumer Name": "Known Consumer"})],
    )
    assert result.records[0]["consumer_number"] == "C-TRACE-01"
    assert result.records[0]["consumer_name"] == "Known Consumer"


# ── Duplicate consumer keys ──────────────────────────────────────────────────


def test_duplicate_consumer_keys_detected():
    rows = [
        _base_source_row(**{"Consumer Number": "DUP001", "Account Number": "A1"}),
        _base_source_row(**{"Consumer Number": "DUP001", "Account Number": "A2"}),
    ]
    result = _run("WC", "HTCT", rows)
    dup_entries = [e for e in result.audit.entries if e.action == "DUPLICATE"]
    assert len(dup_entries) >= 1
    assert any(e.status == AuditStatus.FAILED.value for e in result.audit.entries)


# ── Empty source dataset ────────────────────────────────────────────────────


def test_empty_source_dataset():
    result = _run("WC", "HTCT", [])
    assert len(result.records) == 0
    assert not result.ready_for_output


# ── Missing mandatory source columns ────────────────────────────────────────


def test_missing_mandatory_source_columns():
    rows = [{"Foo": "Bar", "Baz": "Qux"}]
    result = _run("WC", "HTCT", rows)
    # Should have validation warnings about missing columns
    warnings = [w for w in result.validation.warnings if "column" in w.message.lower()]
    assert len(warnings) > 0


# ── Extra source columns ────────────────────────────────────────────────────


def test_extra_source_columns_ignored():
    row = _base_source_row(**{"Extra Custom Column": "extra_value", "Another Extra": "val2"})
    result = _run("WC", "HTCT", [row])
    assert len(result.records) == 1
    rec = result.records[0]
    # Extra columns must not appear in output
    assert "Extra Custom Column" not in rec
    assert "Another Extra" not in rec


# ── Output structure ─────────────────────────────────────────────────────────


class TestOutputStructure:
    def test_output_column_order_matches_template(self):
        result = _run("WC", "HTCT", [_base_source_row()])
        rec = result.records[0]
        output_keys = list(rec.keys())
        assert output_keys == result.template_schema.headers

    def test_output_preserves_all_template_headers(self):
        result = _run("WC", "HTCT", [_base_source_row()])
        rec = result.records[0]
        for header in result.template_schema.headers:
            assert header in rec, f"Template header '{header}' missing from output"

    def test_no_extra_columns_in_output(self):
        result = _run("WC", "HTCT", [_base_source_row()])
        rec = result.records[0]
        template_set = set(result.template_schema.headers)
        output_set = set(rec.keys())
        extra = output_set - template_set
        assert not extra, f"Extra columns in output: {extra}"


# ── Excel writer output ─────────────────────────────────────────────────────


def test_excel_writer_preserves_reference_sheet():
    d = _tmpdir()
    tpl = _template_path(d)
    result = _run("WC", "HTCT", [_base_source_row()], tmpdir=d)
    out = os.path.join(d, "output.xlsx")
    write_output_workbook(
        template_path=tpl,
        target_sheet="HTCT Consumer Master",
        output_records=result.records,
        output_path=out,
        target_headers=result.template_schema.headers,
    )
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert "ref_project_masters_lov" in wb.sheetnames
    assert "HTCT Consumer Master" in wb.sheetnames
    ws = wb["HTCT Consumer Master"]
    # Check that headers are written
    first_row = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    assert "consumer_number" in first_row
    wb.close()


def test_excel_writer_writes_correct_record_count():
    d = _tmpdir()
    tpl = _template_path(d)
    rows = [
        _base_source_row(**{"Consumer Number": f"C{i:03d}", "Account Number": f"A{i:03d}"})
        for i in range(5)
    ]
    result = _run("WC", "HTCT", rows, tmpdir=d)
    out = os.path.join(d, "output5.xlsx")
    write_output_workbook(
        template_path=tpl,
        target_sheet="HTCT Consumer Master",
        output_records=result.records,
        output_path=out,
        target_headers=result.template_schema.headers,
    )
    from openpyxl import load_workbook
    wb = load_workbook(out)
    ws = wb["HTCT Consumer Master"]
    # Row 1 = headers, Row 2 = type hints, Row 3+ = data
    data_rows = 0
    for row in range(3, ws.max_row + 1):
        if ws.cell(row=row, column=1).value is not None or ws.cell(row=row, column=18).value is not None:
            data_rows += 1
    assert data_rows == 5
    wb.close()


# ── Audit report ─────────────────────────────────────────────────────────────


class TestAuditReport:
    def test_audit_has_entries_for_every_source_row(self):
        rows = [
            _base_source_row(**{"Consumer Number": f"C{i:03d}", "Account Number": f"A{i:03d}"})
            for i in range(3)
        ]
        result = _run("WC", "HTCT", rows)
        assert result.audit.total_records == 3

    def test_failed_lookup_not_marked_success(self):
        result = _run("WC", "HTCT", [_base_source_row(**{"Consumer Identification": "UNKNOWN_IDENT"})])
        entry = result.audit.entries[0]
        assert entry.status != AuditStatus.SUCCESS.value
        assert entry.status in {AuditStatus.FAILED.value, AuditStatus.REQUIRES_REVIEW.value,
                                 AuditStatus.WARNING.value}

    def test_audit_export_to_dataframe(self):
        result = _run("WC", "HTCT", [_base_source_row()])
        df = result.audit.to_dataframe()
        assert len(df) == 1
        assert "Status" in df.columns
        assert "Business Key" in df.columns

    def test_error_report_filters_failures(self):
        rows = [
            _base_source_row(**{"Consumer Number": "OK001", "Account Number": "A001"}),
            _base_source_row(**{
                "Consumer Number": "FAIL001", "Account Number": "A002",
                "Consumer Identification": "NONEXISTENT_VALUE",
            }),
        ]
        result = _run("WC", "HTCT", rows)
        err_df = result.audit.get_error_report()
        # At least the failed record should appear
        assert len(err_df) >= 1


# ── Validation ───────────────────────────────────────────────────────────────


def test_validation_reports_missing_reference_categories():
    """If the template reference sheet lacks a category, validation warns."""
    result = _run("WC", "HTCT", [_base_source_row()])
    # CONNECTION CATEGORY is not in our fixture ref data
    # Check that it's at least warned about
    cat_warnings = [
        w for w in result.validation.warnings
        if "CONNECTION CATEGORY" in w.message.upper()
        or "connection_category" in w.message.lower()
    ]
    assert len(cat_warnings) >= 1


# ── Multiple records ─────────────────────────────────────────────────────────


def test_multiple_records_processed():
    rows = [
        _base_source_row(**{
            "Consumer Number": f"C{i:03d}",
            "Account Number": f"A{i:03d}",
            "Consumer Identification": "ESC" if i % 2 == 0 else "TOD",
        })
        for i in range(10)
    ]
    result = _run("WC", "HTCT", rows)
    assert len(result.records) == 10
    assert result.audit.total_records == 10
    # ESC records should become TOD, TOD records stay TOD
    for i, rec in enumerate(result.records):
        ident = rec.get("consumer_identy", "")
        assert ident.strip() == f"{{{TOD_ID}}}"


# ── meter_category always Main Meter ─────────────────────────────────────────


def _run_with_template(op_src: str, op_dst: str, source_rows: list[dict], template_path: str, tmpdir: str):
    hier = _hierarchy_csv(tmpdir)
    src = _source_excel(tmpdir, source_rows)
    op = operation(op_src, op_dst)
    return run_pipeline(
        template_path=template_path,
        target_sheet="HTCT Consumer Master",
        operation=op,
        hierarchy_path=hier,
        source_path=src,
        source_sheet="Sheet1",
    )


class TestMeterCategoryAlwaysMainMeter:
    """meter_category must always resolve to the Main Meter UUID from the reference sheet."""

    @pytest.mark.parametrize("op_id", list(range(1, 10)))
    def test_all_nine_operations_use_main_meter_uuid(self, op_id):
        op = OPERATIONS[op_id]
        ident = "ESC" if op.source == Category.WC else "TOD"
        result = _run(
            op.source.value,
            op.destination.value,
            [_base_source_row(**{"Consumer Identification": ident})],
        )
        mc = result.records[0].get("meter_category", "").strip()
        assert mc == f"{{{MAIN_METER_ID}}}"
        assert mc.startswith("{") and mc.endswith("}")
        assert "{{" not in mc

    def test_source_category_does_not_override_main_meter(self):
        result = _run(
            "WC",
            "HTCT",
            [_base_source_row(**{"Meter Category": "Check Meter"})],
        )
        mc = result.records[0].get("meter_category", "").strip()
        assert mc == f"{{{MAIN_METER_ID}}}"
        assert "Check Meter" not in mc
        entry = result.audit.entries[0]
        assert not any("Check Meter" in (err or "") for err in entry.errors)

    def test_empty_source_meter_category_still_resolves_main_meter(self):
        result = _run(
            "LTCT",
            "LTCT",
            [_base_source_row(**{"Meter Category": ""})],
        )
        mc = result.records[0].get("meter_category", "").strip()
        assert mc == f"{{{MAIN_METER_ID}}}"

    def test_missing_main_meter_reference_is_validation_error(self):
        d = _tmpdir()
        tpl = _template_path(d)
        from openpyxl import load_workbook
        wb = load_workbook(tpl)
        ref = wb["ref_project_masters_lov"]
        rows_to_delete = []
        for idx, row in enumerate(ref.iter_rows(min_row=2, values_only=False), start=2):
            name = str(row[1].value or "").strip()
            master = str(row[3].value or "").strip()
            if name.upper() == "MAIN METER" and master.upper() == "METER CATEGORY":
                rows_to_delete.append(idx)
        for idx in reversed(rows_to_delete):
            ref.delete_rows(idx)
        wb.save(tpl)
        wb.close()

        result = _run_with_template(
            "WC",
            "HTCT",
            [_base_source_row()],
            tpl,
            d,
        )
        entry = result.audit.entries[0]
        assert entry.errors
        assert any("Main Meter" in (err or "") for err in entry.errors)
        assert result.records[0].get("meter_category", "").strip() == ""
        assert not result.ready_for_output

    def test_invalid_main_meter_uuid_is_validation_error(self):
        d = _tmpdir()
        tpl = _template_path(d)
        from openpyxl import load_workbook
        wb = load_workbook(tpl)
        ref = wb["ref_project_masters_lov"]
        for row in ref.iter_rows(min_row=2):
            name = str(row[1].value or "").strip()
            master = str(row[3].value or "").strip()
            if name.upper() == "MAIN METER" and master.upper() == "METER CATEGORY":
                row[0].value = "not-a-valid-uuid"
                break
        wb.save(tpl)
        wb.close()

        result = _run_with_template(
            "HTCT",
            "WC",
            [_base_source_row()],
            tpl,
            d,
        )
        entry = result.audit.entries[0]
        assert entry.errors
        assert any("invalid" in (err or "").lower() or "Main Meter" in (err or "") for err in entry.errors)
        assert "not-a-valid-uuid" not in result.records[0].get("meter_category", "")
        assert not result.ready_for_output


# ── is_old_mtr_available regression (canonical vs alias category) ────────────


class TestIsOldMtrAvailableCanonicalLookup:
    """Regression: source 'Yes' for is_old_mtr_available must resolve to the
    canonical 'IS OLD METER AVAILABLE' UUID even when the template's reference
    sheet also contains a misspelled 'IS OLD METER AVAILABE' with a different
    UUID.  The old code combined both categories, making 'Yes' ambiguous and
    silently writing a blank."""

    OLD_MTR_YES_UUID = "9a2b42b2-44d3-4f25-8d1d-432b316f0bb3"
    MISSPELLED_YES_UUID = "deadbeef-dead-beef-dead-beefdeadbeef"

    def _make_template_with_misspelled_category(self, tmpdir: str) -> str:
        """Template containing both canonical and misspelled categories."""
        path = os.path.join(tmpdir, "template_misspell.xlsx")
        make_template_workbook(path)
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ref = wb["ref_project_masters_lov"]
        # Add misspelled category rows with a DIFFERENT UUID for "Yes"
        ref.append([self.MISSPELLED_YES_UUID, "Yes", "1", "IS OLD METER AVAILABE"])
        ref.append(["abcdef01-2345-6789-abcd-ef0123456789", "No", "2", "IS OLD METER AVAILABE"])
        wb.save(path)
        wb.close()
        return path

    def test_source_yes_resolves_to_canonical_uuid(self):
        """Source 'Yes' must produce the canonical UUID, not blank."""
        d = _tmpdir()
        tpl = self._make_template_with_misspelled_category(d)
        result = _run_with_template(
            "WC", "HTCT",
            [_base_source_row(**{"Is Old Meter Available?": "Yes"})],
            tpl, d,
        )
        rec = result.records[0]
        output = rec.get("is_old_mtr_available", "").strip()
        # Must NOT be blank
        assert output, (
            "is_old_mtr_available is blank; the lookup was probably treated as "
            "ambiguous across canonical and misspelled categories"
        )
        # Must resolve to the canonical category's UUID
        assert output == f"{{{self.OLD_MTR_YES_UUID}}}", (
            f"Expected canonical UUID {{{self.OLD_MTR_YES_UUID}}}, got {output}"
        )

    def test_source_no_resolves_to_canonical_uuid(self):
        """Source 'No' also uses the canonical category."""
        d = _tmpdir()
        tpl = self._make_template_with_misspelled_category(d)
        result = _run_with_template(
            "WC", "HTCT",
            [_base_source_row(**{"Is Old Meter Available?": "No"})],
            tpl, d,
        )
        rec = result.records[0]
        output = rec.get("is_old_mtr_available", "").strip()
        assert output, "is_old_mtr_available is blank for 'No'"
        from tests.conftest import IS_OLD_MTR_NO_ID
        assert output == f"{{{IS_OLD_MTR_NO_ID}}}"

    def test_canonical_only_template_yes_resolves(self):
        """With no misspelled category at all, 'Yes' still resolves."""
        result = _run("WC", "HTCT", [_base_source_row(**{"Is Old Meter Available?": "Yes"})])
        rec = result.records[0]
        output = rec.get("is_old_mtr_available", "").strip()
        assert output == f"{{{self.OLD_MTR_YES_UUID}}}"

    def test_alias_value_y_resolves_to_yes(self):
        """VALUE_ALIASES maps 'Y' → 'Yes', which should then resolve."""
        result = _run("WC", "HTCT", [_base_source_row(**{"Is Old Meter Available?": "Y"})])
        rec = result.records[0]
        output = rec.get("is_old_mtr_available", "").strip()
        assert output == f"{{{self.OLD_MTR_YES_UUID}}}", (
            f"Value alias 'Y' → 'Yes' should resolve; got {output}"
        )

    def test_meter_category_still_main_meter(self):
        """Unrelated rule: meter_category must still be Main Meter."""
        d = _tmpdir()
        tpl = self._make_template_with_misspelled_category(d)
        result = _run_with_template(
            "WC", "HTCT",
            [_base_source_row(**{"Is Old Meter Available?": "Yes"})],
            tpl, d,
        )
        rec = result.records[0]
        assert rec.get("meter_category", "").strip() == f"{{{MAIN_METER_ID}}}"

    def test_consumer_identy_esc_to_tod_preserved(self):
        """Unrelated rule: ESC → TOD consumer identification is unchanged."""
        d = _tmpdir()
        tpl = self._make_template_with_misspelled_category(d)
        result = _run_with_template(
            "WC", "HTCT",
            [_base_source_row(**{
                "Is Old Meter Available?": "Yes",
                "Consumer Identification": "ESC",
            })],
            tpl, d,
        )
        rec = result.records[0]
        assert rec.get("consumer_identy", "").strip() == f"{{{TOD_ID}}}"
