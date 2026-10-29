"""Write transformed records into a copy of the uploaded template."""

from __future__ import annotations

import logging
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

import openpyxl
from openpyxl.utils import get_column_letter

from .config import Category, DATA_START_ROW, TEXT_IDENTIFIER_FIELDS, UUID_FIELDS

logger = logging.getLogger(__name__)


def generate_output_filename(
    dest_category: Category,
    base_name: str = "Consumer_Master",
    output_dir: str = ".",
    source_category: Optional[Category] = None,
    timestamp: Optional[str] = None,
) -> str:
    stamp = timestamp or datetime.now().strftime("%Y%m%d_%H%M%S")
    if source_category is not None:
        filename = f"{base_name}_{source_category.value}_to_{dest_category.value}_{stamp}.xlsx"
    else:
        filename = f"{dest_category.value}_{base_name}_Updated_{stamp}.xlsx"
    return str(Path(output_dir) / filename)


def write_output_workbook(
    template_path: str,
    target_sheet: str,
    output_records: List[Dict[str, str]],
    output_path: str,
    target_headers: List[str],
    preserve_type_hint_row: bool = True,
) -> str:
    wb = openpyxl.load_workbook(template_path)
    if target_sheet not in wb.sheetnames:
        raise ValueError(f"Target sheet '{target_sheet}' not found in template")
    ws = wb[target_sheet]

    for col_idx, header in enumerate(target_headers, 1):
        existing = ws.cell(row=1, column=col_idx).value
        if existing is None or str(existing).strip() != header:
            ws.cell(row=1, column=col_idx).value = header

    data_start = DATA_START_ROW if preserve_type_hint_row else 2
    for col_idx, header in enumerate(target_headers, 1):
        if header.lower().strip() in TEXT_IDENTIFIER_FIELDS:
            column = ws.column_dimensions[get_column_letter(col_idx)]
            column.width = max(column.width or 0, 20)

    if ws.max_row >= data_start:
        for row in range(data_start, ws.max_row + 1):
            for col in range(1, len(target_headers) + 1):
                ws.cell(row=row, column=col).value = None

    for row_idx, record in enumerate(output_records, data_start):
        for col_idx, header in enumerate(target_headers, 1):
            value = record.get(header, "")
            cell = ws.cell(row=row_idx, column=col_idx)
            header_lower = header.lower().strip()
            if value is None or (isinstance(value, str) and not str(value).strip()):
                cell.value = None
                continue
            text = str(value).strip()
            if header_lower in UUID_FIELDS or header_lower in TEXT_IDENTIFIER_FIELDS:
                cell.value = text
                cell.number_format = "@"
            else:
                try:
                    float_val = float(text)
                    if float_val == int(float_val) and "." not in text:
                        cell.value = int(float_val)
                    else:
                        cell.value = float_val
                except (ValueError, TypeError):
                    cell.value = text

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    wb.close()
    logger.info("Output workbook written: %s (%s records)", output_path, len(output_records))
    return output_path
